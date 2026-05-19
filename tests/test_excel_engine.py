import os
import shutil
import sys
import tempfile
import unittest

import openpyxl
from openpyxl.styles import Font, PatternFill

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from services.excel_export_service import ExcelExportService
from services.excel_template_engine import ExcelTemplateEngine, ExcelTemplateError


class TestExcelTemplateEngine(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()
        self.engine = ExcelTemplateEngine()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def _path(self, name):
        return os.path.join(self.tmp, name)

    def test_scalar_replacement(self):
        path = self._path("scalar.xlsx")
        out = self._path("scalar_out.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "{{ CourseName }}"
        ws["B1"] = "Code: {{ CourseCode }}"
        wb.save(path)

        self.engine.render(path, {"CourseName": "Math", "CourseCode": "M101"}, out)
        rendered = openpyxl.load_workbook(out)
        ws = rendered.active
        self.assertEqual(ws["A1"].value, "Math")
        self.assertEqual(ws["B1"].value, "Code: M101")

    def test_loop_rendering_and_style_preserved(self):
        path = self._path("loop.xlsx")
        out = self._path("loop_out.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "{{#CLOs}}"
        ws["A2"] = "{{ Code }}"
        ws["B2"] = "{{ Description }}"
        ws["A2"].font = Font(bold=True)
        ws["B2"].fill = PatternFill("solid", fgColor="FFFF00")
        ws.row_dimensions[2].height = 33
        ws["A3"] = "{{/CLOs}}"
        wb.save(path)

        self.engine.render(
            path,
            {"CLOs": [{"Code": "CLO1", "Description": "Analyze"}, {"Code": "CLO2", "Description": "Design"}]},
            out,
        )
        rendered = openpyxl.load_workbook(out)
        ws = rendered.active
        self.assertEqual(ws["A1"].value, "CLO1")
        self.assertEqual(ws["B2"].value, "Design")
        self.assertTrue(ws["A1"].font.bold)
        self.assertEqual(ws["B1"].fill.fgColor.rgb, "00FFFF00")
        self.assertEqual(ws.row_dimensions[1].height, 33)

    def test_invalid_placeholder_has_suggestion(self):
        path = self._path("invalid.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "{{ CourseNam }}"
        wb.save(path)

        result = self.engine.validate_template(path, {})
        self.assertFalse(result["valid"])
        self.assertIn("CourseName", result["errors"][0]["suggestion"])

    def test_merged_cell_preservation_in_loop(self):
        path = self._path("merged.xlsx")
        out = self._path("merged_out.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws["A1"] = "{{#CLOs}}"
        ws["A2"] = "{{ Code }}"
        ws["B2"] = "{{ Description }}"
        ws.merge_cells("B2:C2")
        ws["A3"] = "{{/CLOs}}"
        wb.save(path)

        self.engine.render(
            path,
            {"CLOs": [{"Code": "CLO1", "Description": "A"}, {"Code": "CLO2", "Description": "B"}]},
            out,
        )
        rendered = openpyxl.load_workbook(out)
        ranges = {str(rng) for rng in rendered.active.merged_cells.ranges}
        self.assertIn("B1:C1", ranges)
        self.assertIn("B2:C2", ranges)

    def test_default_template_can_render(self):
        path = self._path("default.xlsx")
        out = self._path("default_out.xlsx")
        self.engine.create_default_template(path)
        self.engine.render(
            path,
            {
                "CourseName": "Math",
                "CourseCode": "M101",
                "Credits": 3,
                "CLOs": [{"Code": "CLO1", "Description": "Analyze", "PLO": "PLO1"}],
                "Rubrics": [{"Code": "RB1", "Name": "Rubric", "Description": "Criteria"}],
                "RevisionHistory": [{"Version": "1", "Description": "Created", "Date": "2026-05-19", "UpdatedBy": "QA"}],
                "SignPlace": "Ha Noi",
                "SignDate": "19/05/2026",
                "SignerLeftTitle": "Author",
                "SignerLeftName": "Nguyen A",
                "SignerRightTitle": "Dean",
                "SignerRightName": "Tran B",
            },
            out,
        )
        self.assertTrue(os.path.exists(out))
        rendered = openpyxl.load_workbook(out)
        self.assertIn("Rubrics", rendered.sheetnames)
        self.assertIn("RevisionHistory", rendered.sheetnames)
        self.assertIn("Approval", rendered.sheetnames)
        self.assertEqual(rendered["Rubrics"]["A4"].value, "RB1")
        self.assertEqual(rendered["RevisionHistory"]["A4"].value, "1")
        self.assertEqual(rendered["Approval"]["B3"].value, "Ha Noi")


class FakeExcelTemplates:
    def __init__(self):
        self.default = {"id": 1}

    def get_default(self):
        return self.default

    def render(self, template_id, context, output_path):
        wb = openpyxl.Workbook()
        wb.active["A1"] = context["CourseCode"]
        wb.save(output_path)


class FakeDB:
    def __init__(self):
        self.logs = []

    def add_import_export_log(self, **kwargs):
        self.logs.append(kwargs)


class TestExcelExportService(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def tearDown(self):
        shutil.rmtree(self.tmp)

    def test_batch_export_continues_and_logs(self):
        svc = ExcelExportService.__new__(ExcelExportService)
        svc.db = FakeDB()
        svc.excel_templates = FakeExcelTemplates()
        svc.build_context = lambda hp_id: {"CourseCode": f"HP{hp_id}"} if hp_id != 2 else (_ for _ in ()).throw(ValueError("bad"))

        result = svc.export_batch([1, 2, 3], self.tmp)
        self.assertEqual(result["success"], 2)
        self.assertEqual(result["errors"], 1)
        self.assertTrue(os.path.exists(os.path.join(self.tmp, "HP1_20260519.xlsx")) or len(os.listdir(self.tmp)) == 2)
        self.assertEqual(svc.db.logs[0]["type"], "EXPORT_EXCEL")


if __name__ == "__main__":
    unittest.main()
