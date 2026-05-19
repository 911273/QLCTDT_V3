"""Dynamic Excel template renderer based on openpyxl."""

from __future__ import annotations

import copy
import difflib
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Any

# Delay importing openpyxl until actually needed to avoid slowing startup
# (openpyxl will be imported lazily inside the class)

from services.placeholder_registry import (
    flatten_registry,
    normalize_template_context,
    valid_child_keys,
    valid_loop_keys,
    valid_scalar_keys,
)


PLACEHOLDER_RE = re.compile(r"\{\{\s*([A-Za-z_][\w.]*)\s*\}\}")
LOOP_START_RE = re.compile(r"^\s*\{\{\#([A-Za-z_][\w]*)\}\}\s*$")
LOOP_END_RE = re.compile(r"^\s*\{\{\/([A-Za-z_][\w]*)\}\}\s*$")


class ExcelTemplateError(Exception):
    pass


@dataclass
class TemplateIssue:
    sheet: str
    cell: str
    message: str
    suggestion: str = ""

    def as_dict(self) -> dict:
        data = {"sheet": self.sheet, "cell": self.cell, "message": self.message}
        if self.suggestion:
            data["suggestion"] = self.suggestion
        return data


@dataclass
class LoopBlock:
    name: str
    start_row: int
    end_row: int


class ExcelTemplateEngine:
    """Render .xlsx templates with scalar placeholders and row loop blocks."""

    def __init__(self) -> None:
        self._openpyxl_loaded = False

    def _import_openpyxl(self) -> None:
        if self._openpyxl_loaded:
            return
        try:
            from openpyxl import load_workbook, Workbook
            from openpyxl.cell.cell import MergedCell
            from openpyxl.styles import Alignment, Border, Font, PatternFill
            from openpyxl.utils import get_column_letter
        except Exception as exc:
            raise ExcelTemplateError(f"Missing openpyxl dependency: {exc}") from exc

        self.load_workbook = load_workbook
        self.Workbook = Workbook
        self.MergedCell = MergedCell
        self.Alignment = Alignment
        self.Border = Border
        self.Font = Font
        self.PatternFill = PatternFill
        self.get_column_letter = get_column_letter
        self._openpyxl_loaded = True

    def render(self, template_path: str, context: dict, output_path: str) -> None:
        self._import_openpyxl()
        validation = self.validate_template(template_path, context)
        if not validation["valid"]:
            lines = [
                f"Sheet: {i['sheet']} Cell: {i['cell']} {i['message']}"
                + (f" Suggestion: {i['suggestion']}" if i.get("suggestion") else "")
                for i in validation["errors"]
            ]
            raise ExcelTemplateError("Invalid Excel template:\n" + "\n".join(lines))

        normalized = normalize_template_context(context)
        wb = self.load_workbook(template_path)

        for ws in wb.worksheets:
            self._render_sheet_loops(ws, normalized)
            self._render_sheet_scalars(ws, normalized)

        out = Path(output_path)
        if out.parent:
            out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)

    def scan_placeholders(self, path: str) -> dict:
        self._import_openpyxl()
        wb = self.load_workbook(path, data_only=False)
        scalars = []
        loops = []
        for ws in wb.worksheets:
            for row in ws.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    text = cell.value.strip()
                    start = LOOP_START_RE.match(text)
                    end = LOOP_END_RE.match(text)
                    if start:
                        loops.append({"sheet": ws.title, "cell": cell.coordinate, "name": start.group(1), "kind": "start"})
                        continue
                    if end:
                        loops.append({"sheet": ws.title, "cell": cell.coordinate, "name": end.group(1), "kind": "end"})
                        continue
                    for name in PLACEHOLDER_RE.findall(cell.value):
                        scalars.append({"sheet": ws.title, "cell": cell.coordinate, "name": name})
        return {"scalars": scalars, "loops": loops}

    def validate_template(self, path: str, context: dict | None = None) -> dict:
        self._import_openpyxl()
        try:
            wb = self.load_workbook(path, data_only=False)
        except Exception as exc:
            raise ExcelTemplateError(f"Cannot open workbook: {exc}") from exc

        context = normalize_template_context(context or {})
        scalar_keys = set(context.keys()) | valid_scalar_keys()
        loop_keys = valid_loop_keys()
        issues: list[TemplateIssue] = []
        scalars = []
        loops = []

        for ws in wb.worksheets:
            stack: list[tuple[str, str, int]] = []
            for row in ws.iter_rows():
                for cell in row:
                    if not isinstance(cell.value, str):
                        continue
                    text = cell.value.strip()
                    start = LOOP_START_RE.match(text)
                    end = LOOP_END_RE.match(text)
                    if start:
                        name = start.group(1)
                        loops.append({"sheet": ws.title, "cell": cell.coordinate, "name": name, "kind": "start"})
                        if name not in loop_keys and name not in context:
                            issues.append(self._unknown(ws.title, cell.coordinate, name, loop_keys | set(context.keys()), "Unknown loop"))
                        stack.append((name, cell.coordinate, cell.row))
                        continue
                    if end:
                        name = end.group(1)
                        loops.append({"sheet": ws.title, "cell": cell.coordinate, "name": name, "kind": "end"})
                        if not stack:
                            issues.append(TemplateIssue(ws.title, cell.coordinate, f"Unexpected loop end: {{{{/ {name} }}}}".replace("/ ", "/")))
                        else:
                            open_name, open_cell, open_row = stack.pop()
                            if open_name != name:
                                issues.append(
                                    TemplateIssue(
                                        ws.title,
                                        cell.coordinate,
                                        f"Loop end does not match start at {open_cell}: {name} != {open_name}",
                                    )
                                )
                            elif cell.row <= open_row + 1:
                                issues.append(TemplateIssue(ws.title, cell.coordinate, f"Loop {name} has no template row"))
                        continue

                    loop_name = self._active_loop_for_cell(ws, cell.row)
                    for name in PLACEHOLDER_RE.findall(cell.value):
                        scalars.append({"sheet": ws.title, "cell": cell.coordinate, "name": name})
                        if loop_name:
                            allowed = valid_child_keys(loop_name)
                            if allowed and name not in allowed:
                                issues.append(self._unknown(ws.title, cell.coordinate, name, allowed, "Unknown loop field"))
                        elif name not in scalar_keys:
                            issues.append(self._unknown(ws.title, cell.coordinate, name, scalar_keys, "Unknown placeholder"))

            for name, cell_ref, _row in stack:
                issues.append(TemplateIssue(ws.title, cell_ref, f"Unclosed loop: {name}"))

        return {
            "valid": not issues,
            "errors": [issue.as_dict() for issue in issues],
            "scalars": scalars,
            "loops": loops,
            "total_keys": len(scalars) + len(loops),
        }

    def create_default_template(self, output_path: str) -> str:
        self._import_openpyxl()
        wb = self.Workbook()
        ws = wb.active
        ws.title = "Overview"
        self._write_overview(ws)
        self._write_loop_sheet(wb.create_sheet("CLOs"), "CLOs", ["Code", "Description", "PLO", "Level"])
        self._write_loop_sheet(wb.create_sheet("Objectives"), "Objectives", ["No", "Description", "PLO"])
        self._write_loop_sheet(
            wb.create_sheet("TeachingContents"),
            "TeachingContents",
            ["No", "Title", "HoursLT", "HoursBT", "HoursTL", "HoursTH", "TeachingMethod", "LearningTask", "CLO", "Assessment"],
        )
        self._write_loop_sheet(
            wb.create_sheet("Assessment"),
            "AssessmentRows",
            ["No", "Group", "Content", "Method", "Time", "Scale", "Weight", "CLO"],
        )
        self._write_loop_sheet(wb.create_sheet("Rubrics"), "Rubrics", ["Code", "Name", "Description"])
        self._write_loop_sheet(wb.create_sheet("Materials"), "MainMaterials", ["No", "Title", "Author", "Info", "Content"])
        self._write_loop_sheet(wb.create_sheet("References"), "ReferenceMaterials", ["No", "Title", "Author", "Info", "Content"])
        self._write_loop_sheet(wb.create_sheet("RevisionHistory"), "RevisionHistory", ["Version", "Description", "Date", "UpdatedBy"])
        self._write_approval_sheet(wb.create_sheet("Approval"))
        self._write_loop_sheet(wb.create_sheet("Policies"), "Policies", ["Type", "Content"])
        self._write_loop_sheet(wb.create_sheet("Checklist"), "Checklists", ["Item", "Status"])

        out = Path(output_path)
        if out.parent:
            out.parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path

    def create_placeholder_catalog(self, output_path: str) -> str:
        self._import_openpyxl()
        wb = self.Workbook()
        ws = wb.active
        ws.title = "Scalars"
        headers = ["Placeholder", "Type", "Group", "Description"]
        self._write_headers(ws, headers)
        row_no = 2
        loops = wb.create_sheet("Loops")
        self._write_headers(loops, ["Loop", "Child fields", "Example"])
        loop_row = 2
        for row in flatten_registry():
            if row["type"] == "scalar":
                ws.append([f"{{{{ {row['key']} }}}}", row["type"], row["group"], row.get("desc", "")])
            elif row["type"] == "loop":
                children = ", ".join(row.get("children", []))
                example = f"{{{{#{row['key']}}}}} ... {{{{/{row['key']}}}}}"
                loops.append([row["key"], children, example])
                loop_row += 1
            row_no += 1
        for sheet in [ws, loops]:
            for col in range(1, sheet.max_column + 1):
                sheet.column_dimensions[self.get_column_letter(col)].width = 28
        Path(output_path).parent.mkdir(parents=True, exist_ok=True)
        wb.save(output_path)
        return output_path

    def _render_sheet_loops(self, ws: Worksheet, context: dict) -> None:
        blocks = self._find_loop_blocks(ws)
        for block in sorted(blocks, key=lambda b: b.start_row, reverse=True):
            items = context.get(block.name) or []
            if not isinstance(items, list):
                raise ExcelTemplateError(f"Loop '{block.name}' is not a list")
            self._render_loop_block(ws, block, items)

    def _render_sheet_scalars(self, ws: Worksheet, context: dict) -> None:
        self._import_openpyxl()
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell, self.MergedCell) or not isinstance(cell.value, str):
                    continue
                cell.value = self._render_text(cell.value, context)

    def _render_loop_block(self, ws: Worksheet, block: LoopBlock, items: list[dict]) -> None:
        template_start = block.start_row + 1
        template_end = block.end_row - 1
        template_height = template_end - template_start + 1
        if template_height <= 0:
            raise ExcelTemplateError(f"Loop '{block.name}' has no template row")

        max_col = ws.max_column
        rows_snapshot = []
        for row_idx in range(template_start, template_end + 1):
            cells = []
            for col_idx in range(1, max_col + 1):
                cell = ws.cell(row_idx, col_idx)
                cells.append(self._snapshot_cell(cell))
            rows_snapshot.append({"height": ws.row_dimensions[row_idx].height, "cells": cells})

        merge_ranges = self._snapshot_merges(ws, template_start, template_end)
        ws.delete_rows(block.start_row, block.end_row - block.start_row + 1)

        insert_count = len(items) * template_height
        if insert_count:
            ws.insert_rows(block.start_row, insert_count)

        current_row = block.start_row
        for item in items:
            item_context = normalize_template_context(item)
            item_context.update(item)
            for row_template in rows_snapshot:
                ws.row_dimensions[current_row].height = row_template["height"]
                for col_idx, snapshot in enumerate(row_template["cells"], start=1):
                    self._apply_snapshot(ws.cell(current_row, col_idx), snapshot, item_context)
                current_row += 1
            self._restore_merges(ws, merge_ranges, block.start_row, current_row - template_height)

    def _find_loop_blocks(self, ws: Worksheet) -> list[LoopBlock]:
        starts = {}
        blocks = []
        for row in ws.iter_rows():
            for cell in row:
                if not isinstance(cell.value, str):
                    continue
                start = LOOP_START_RE.match(cell.value.strip())
                end = LOOP_END_RE.match(cell.value.strip())
                if start:
                    starts[start.group(1)] = cell.row
                elif end:
                    name = end.group(1)
                    if name in starts:
                        blocks.append(LoopBlock(name=name, start_row=starts.pop(name), end_row=cell.row))
        return blocks

    def _active_loop_for_cell(self, ws: Worksheet, row_idx: int) -> str | None:
        for block in self._find_loop_blocks(ws):
            if block.start_row < row_idx < block.end_row:
                return block.name
        return None

    def _render_text(self, text: str, context: dict) -> Any:
        matches = PLACEHOLDER_RE.findall(text)
        if not matches:
            return text
        if len(matches) == 1 and text.strip() == f"{{{{ {matches[0]} }}}}":
            return context.get(matches[0], "")
        result = text
        for name in matches:
            result = re.sub(r"\{\{\s*" + re.escape(name) + r"\s*\}\}", str(context.get(name, "")), result)
        return result

    def _snapshot_cell(self, cell) -> dict:
        return {
            "value": cell.value,
            "style": copy.copy(cell._style),
            "font": copy.copy(cell.font),
            "fill": copy.copy(cell.fill),
            "border": copy.copy(cell.border),
            "alignment": copy.copy(cell.alignment),
            "number_format": cell.number_format,
            "protection": copy.copy(cell.protection),
        }

    def _apply_snapshot(self, cell, snapshot: dict, context: dict) -> None:
        cell._style = copy.copy(snapshot["style"])
        cell.font = copy.copy(snapshot["font"])
        cell.fill = copy.copy(snapshot["fill"])
        cell.border = copy.copy(snapshot["border"])
        cell.alignment = copy.copy(snapshot["alignment"])
        cell.number_format = snapshot["number_format"]
        cell.protection = copy.copy(snapshot["protection"])
        value = snapshot["value"]
        cell.value = self._render_text(value, context) if isinstance(value, str) else value

    def _snapshot_merges(self, ws: Worksheet, start_row: int, end_row: int) -> list[tuple[int, int, int, int]]:
        ranges = []
        to_unmerge = []
        for merged in list(ws.merged_cells.ranges):
            if start_row <= merged.min_row and merged.max_row <= end_row:
                ranges.append((merged.min_row - start_row, merged.max_row - start_row, merged.min_col, merged.max_col))
                to_unmerge.append(str(merged))
        for ref in to_unmerge:
            ws.unmerge_cells(ref)
        return ranges

    def _restore_merges(self, ws: Worksheet, merge_ranges: list[tuple[int, int, int, int]], template_start: int, target_start: int) -> None:
        for min_off, max_off, min_col, max_col in merge_ranges:
            ws.merge_cells(
                start_row=target_start + min_off,
                end_row=target_start + max_off,
                start_column=min_col,
                end_column=max_col,
            )

    def _unknown(self, sheet: str, cell: str, name: str, valid_names: set[str], prefix: str) -> TemplateIssue:
        suggestion = ""
        match = difflib.get_close_matches(name, sorted(valid_names), n=1)
        if match:
            suggestion = f"{{{{ {match[0]} }}}}"
        return TemplateIssue(sheet, cell, f"{prefix}: {{{{ {name} }}}}", suggestion)

    def _write_overview(self, ws: Worksheet) -> None:
        self._import_openpyxl()
        ws["A1"] = "DE CUONG CHI TIET HOC PHAN"
        ws["A1"].font = self.Font(bold=True, size=14, color="FFFFFF")
        ws["A1"].fill = self.PatternFill("solid", fgColor="1F4E79")
        ws.merge_cells("A1:D1")
        rows = [
            ("Ten hoc phan", "{{ CourseName }}"),
            ("Ma hoc phan", "{{ CourseCode }}"),
            ("So tin chi", "{{ Credits }}"),
            ("Don vi", "{{ Department }}"),
            ("Giang vien phu trach", "{{ Lecturer }}"),
            ("Loai hoc phan", "{{ CourseType }}"),
            ("Hoc ky", "{{ Semester }}"),
            ("Chuong trinh", "{{ ProgramName }}"),
            ("Nganh/chuyen nganh", "{{ MajorName }}"),
            ("Tom tat", "{{ Summary }}"),
            ("Phuong phap giang day", "{{ TeachingMethod }}"),
            ("Hoc phan tien quyet", "{{ Prerequisite }}"),
            ("Hoc phan thay the", "{{ ReplacementCourse }}"),
        ]
        for row_idx, (label, placeholder) in enumerate(rows, start=3):
            ws.cell(row_idx, 1, label)
            ws.cell(row_idx, 2, placeholder)
            ws.cell(row_idx, 1).font = self.Font(bold=True)
            ws.cell(row_idx, 2).alignment = self.Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 60

    def _write_approval_sheet(self, ws: Worksheet) -> None:
        self._import_openpyxl()
        ws["A1"] = "Approval"
        ws["A1"].font = self.Font(bold=True, size=13, color="FFFFFF")
        ws["A1"].fill = self.PatternFill("solid", fgColor="1F4E79")
        ws.merge_cells("A1:D1")
        rows = [
            ("Dia diem ky", "{{ SignPlace }}"),
            ("Ngay ky", "{{ SignDate }}"),
            ("Nguoi lap - chuc danh", "{{ SignerLeftTitle }}"),
            ("Nguoi lap - ho ten", "{{ SignerLeftName }}"),
            ("Nguoi duyet - chuc danh", "{{ SignerRightTitle }}"),
            ("Nguoi duyet - ho ten", "{{ SignerRightName }}"),
        ]
        for row_idx, (label, placeholder) in enumerate(rows, start=3):
            ws.cell(row_idx, 1, label)
            ws.cell(row_idx, 2, placeholder)
            ws.cell(row_idx, 1).font = self.Font(bold=True)
            ws.cell(row_idx, 2).alignment = self.Alignment(wrap_text=True)
        ws.column_dimensions["A"].width = 28
        ws.column_dimensions["B"].width = 48

    def _write_loop_sheet(self, ws: Worksheet, loop_name: str, fields: list[str]) -> None:
        self._import_openpyxl()
        ws["A1"] = loop_name
        ws["A1"].font = self.Font(bold=True, size=13, color="FFFFFF")
        ws["A1"].fill = self.PatternFill("solid", fgColor="1F4E79")
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(fields)))
        self._write_headers(ws, fields, row=3)
        ws.cell(4, 1, f"{{{{#{loop_name}}}}}")
        for col_idx, field in enumerate(fields, start=1):
            cell = ws.cell(5, col_idx, f"{{{{ {field} }}}}")
            cell.alignment = self.Alignment(wrap_text=True, vertical="top")
            cell.border = self.Border()
        ws.cell(6, 1, f"{{{{/{loop_name}}}}}")
        for col_idx in range(1, len(fields) + 1):
            ws.column_dimensions[self.get_column_letter(col_idx)].width = 18 if col_idx != 2 else 50

    def _write_headers(self, ws: Worksheet, headers: list[str], row: int = 1) -> None:
        self._import_openpyxl()
        for col_idx, header in enumerate(headers, start=1):
            cell = ws.cell(row, col_idx, header)
            cell.font = self.Font(bold=True, color="FFFFFF")
            cell.fill = self.PatternFill("solid", fgColor="5B9BD5")
            cell.alignment = self.Alignment(horizontal="center", vertical="center", wrap_text=True)
