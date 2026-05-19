# template_manager_dialog.py
"""
Dialog quản lý Word Template v2.
Cho phép người dùng: Upload, Preview, Validate, Đặt mặc định, Xóa template.
Tích hợp với TemplateService (services/template_service.py).
"""

import os
import json
import tkinter as tk
from tkinter import messagebox, filedialog
import ttkbootstrap as tb
from ttkbootstrap.constants import *
from ui.theme.design_system import UI_THEME
from ui.widgets.dialog_base import configure_dialog
from ui.widgets.searchable_tree import apply_tree_defaults, restripe_tree


class TemplateManagerDialog(tb.Toplevel):
    """
    Dialog quản lý Template Word linh hoạt.
    
    Sử dụng:
        dlg = TemplateManagerDialog(parent, db, hp_id=current_hp_id)
        parent.wait_window(dlg)
        if dlg.selected_template_id:   # user chọn export với template cụ thể
            ...
    """

    def __init__(self, parent, db, hp_id: int = None):
        super().__init__(parent)
        self.db = db
        self.hp_id = hp_id
        self.selected_template_id = None

        # Import service
        from services.template_service import TemplateService
        self.svc = TemplateService(db)

        # Window setup
        self.title("📋 Quản lý Template Word")
        configure_dialog(self, parent, width=920, height=600, min_width=750, min_height=480)
        self._build_ui()
        self._refresh_list()

    def _center(self):
        self.update_idletasks()
        x = (self.winfo_screenwidth() - 920) // 2
        y = (self.winfo_screenheight() - 600) // 2
        self.geometry(f"+{x}+{y}")

    # ── UI Build ─────────────────────────────────────────────────────────────

    def _build_ui(self):
        self.configure(padx=UI_THEME["padding_x"], pady=UI_THEME["padding_y"])

        # Title bar
        hdr = tb.Frame(self, bootstyle='dark')
        hdr.pack(fill='x', pady=(0, 10))
        tb.Label(hdr, text="  📋 Quản lý Template Word Linh hoạt",
                 font=('Segoe UI', 13, 'bold'), bootstyle='inverse-dark').pack(
                     side='left', fill='x', pady=6)

        # Main split
        split = tb.Frame(self)
        split.pack(fill='both', expand=True)

        # LEFT — danh sách template
        left = tb.Labelframe(split, text=" 📁 Danh sách Template ", bootstyle='primary', padding=8)
        left.pack(side='left', fill='both', padx=(0, 8), pady=4, ipadx=2)

        # Treeview
        self.tree = tb.Treeview(
            left,
            columns=('ten', 'default', 'ngay'),
            show='headings',
            height=14,
            bootstyle='primary'
        )
        apply_tree_defaults(self.tree)
        self.tree.heading('ten', text='Tên Template')
        self.tree.heading('default', text='Mặc định')
        self.tree.heading('ngay', text='Ngày tạo')
        self.tree.column('ten', width=200, anchor='w')
        self.tree.column('default', width=70, anchor='center')
        self.tree.column('ngay', width=95, anchor='center')
        self.tree.pack(fill='both', expand=True)
        self.tree.bind('<<TreeviewSelect>>', self._on_select)

        # Action buttons
        btn_frame = tb.Frame(left)
        btn_frame.pack(fill='x', pady=(6, 0))
        tb.Button(btn_frame, text="➕ Upload .docx", bootstyle='success', width=14,
                  command=self._upload).pack(side='left', padx=2)
        tb.Button(btn_frame, text="🗑 Xóa", bootstyle='danger-outline', width=8,
                  command=self._delete).pack(side='left', padx=2)

        # RIGHT — chi tiết
        self.right = tb.Labelframe(split, text=" 🔍 Chi tiết Template ", bootstyle='info', padding=10)
        self.right.pack(side='left', fill='both', expand=True, pady=4)
        self._build_right()

    def _build_right(self):
        right = self.right

        # Name
        row1 = tb.Frame(right)
        row1.pack(fill='x', pady=3)
        tb.Label(row1, text="Tên:", width=12, anchor='e').pack(side='left')
        self.var_ten = tk.StringVar()
        tb.Entry(row1, textvariable=self.var_ten, width=30).pack(side='left', padx=4)

        # Description
        row2 = tb.Frame(right)
        row2.pack(fill='x', pady=3)
        tb.Label(row2, text="Mô tả:", width=12, anchor='e').pack(side='left')
        self.var_mo_ta = tk.StringVar()
        tb.Entry(row2, textvariable=self.var_mo_ta, width=40).pack(side='left', padx=4)

        # File path
        row3 = tb.Frame(right)
        row3.pack(fill='x', pady=3)
        tb.Label(row3, text="File:", width=12, anchor='e').pack(side='left')
        self.var_path = tk.StringVar()
        tb.Label(row3, textvariable=self.var_path, foreground='gray',
                 font=('Consolas', 8)).pack(side='left', padx=4)

        # Buttons row
        btn_row = tb.Frame(right)
        btn_row.pack(fill='x', pady=(8, 4))
        tb.Button(btn_row, text="⭐ Đặt làm mặc định", bootstyle='warning', width=18,
                  command=self._set_default).pack(side='left', padx=2)
        tb.Button(btn_row, text="✏ Cập nhật tên", bootstyle='secondary-outline', width=14,
                  command=self._save_name).pack(side='left', padx=2)

        # Separator
        tb.Separator(right, orient='horizontal').pack(fill='x', pady=8)

        # Validate section
        tb.Label(right, text="📌 Kiểm tra Placeholder:", font=('Segoe UI', 9, 'bold')).pack(anchor='w')
        tb.Button(right, text="🔍 Validate Template", bootstyle='info-outline',
                  command=self._validate).pack(anchor='w', pady=4)

        self.txt_validate = tk.Text(right, height=6, state='disabled', font=('Consolas', 8),
                                    bg='#1e1e2e', fg='#cdd6f4', relief='flat',
                                    wrap='word')
        self.txt_validate.pack(fill='both', expand=True, pady=2)

        # Separator
        tb.Separator(right, orient='horizontal').pack(fill='x', pady=8)

        # Preview section
        preview_row = tb.Frame(right)
        preview_row.pack(fill='x')
        tb.Label(preview_row, text="👁 Xem trước với HP hiện tại:", font=('Segoe UI', 9, 'bold')).pack(side='left')
        tb.Button(preview_row, text="📄 Preview & Export", bootstyle='primary',
                  command=self._preview_export).pack(side='right')

        # Placeholder help
        tb.Separator(right, orient='horizontal').pack(fill='x', pady=8)
        tb.Button(right, text="📋 Xem danh sách tất cả Placeholder", bootstyle='light',
                  command=self._show_placeholders).pack(anchor='w')

        # Bottom: select & close
        btn_bottom = tb.Frame(right)
        btn_bottom.pack(fill='x', pady=(10, 0))
        if self.hp_id:
            tb.Button(btn_bottom, text="✅ Xuất Word với Template này", bootstyle='success', width=24,
                      command=self._export_with_this).pack(side='left', padx=2)
        tb.Button(btn_bottom, text="✖ Đóng", bootstyle='secondary-outline', width=10,
                  command=self.destroy).pack(side='right', padx=2)

    # ── Data ─────────────────────────────────────────────────────────────────

    def _refresh_list(self):
        for item in self.tree.get_children():
            self.tree.delete(item)

        self._templates = self.svc.get_all()
        for t in self._templates:
            tag = 'default' if t.get('la_mac_dinh') else ''
            self.tree.insert('', 'end',
                             iid=str(t['id']),
                             values=(
                                 t['ten'],
                                 '✅ Có' if t.get('la_mac_dinh') else '-',
                                 (t.get('ngay_tao') or '')[:10]
                              ),
                             tags=(tag,))
        self.tree.tag_configure('default', foreground='#a6e3a1')
        restripe_tree(self.tree)

    def _get_selected_id(self) -> int:
        sel = self.tree.selection()
        return int(sel[0]) if sel else None

    def _get_selected_tpl(self) -> dict:
        tid = self._get_selected_id()
        if not tid:
            return {}
        return next((t for t in self._templates if t['id'] == tid), {})

    def _on_select(self, event=None):
        tpl = self._get_selected_tpl()
        if not tpl:
            return
        self.var_ten.set(tpl.get('ten', ''))
        self.var_mo_ta.set(tpl.get('mo_ta', ''))
        fp = tpl.get('file_path', '')
        self.var_path.set(os.path.basename(fp) if fp else '(Chưa có file)')

    # ── Actions ──────────────────────────────────────────────────────────────

    def _upload(self):
        path = filedialog.askopenfilename(
            title="Chọn file template Word",
            filetypes=[("Word Documents", "*.docx"), ("All Files", "*.*")]
        )
        if not path:
            return

        ten = os.path.splitext(os.path.basename(path))[0]
        # Hỏi tên
        dlg = _TextInputDialog(self, "Tên Template", "Nhập tên template:", ten)
        self.wait_window(dlg)
        if dlg.result is None:
            return
        ten = dlg.result or ten

        try:
            tpl_id = self.svc.upload(path, ten)
            self._refresh_list()
            # Chọn template mới upload
            self.tree.selection_set(str(tpl_id))
            self.tree.see(str(tpl_id))
            self._on_select()
            messagebox.showinfo("Thành công", f"Đã upload template '{ten}' thành công!", parent=self)
        except Exception as e:
            messagebox.showerror("Lỗi upload", str(e), parent=self)

    def _delete(self):
        tid = self._get_selected_id()
        if not tid:
            messagebox.showwarning("Chưa chọn", "Hãy chọn template cần xóa.", parent=self)
            return
        tpl = self._get_selected_tpl()
        if not messagebox.askyesno("Xác nhận xóa", f"Xóa template '{tpl.get('ten', '')}'?", parent=self):
            return
        try:
            self.svc.delete(tid)
            self._refresh_list()
            self._clear_detail()
        except Exception as e:
            messagebox.showerror("Lỗi xóa", str(e), parent=self)

    def _set_default(self):
        tid = self._get_selected_id()
        if not tid:
            messagebox.showwarning("Chưa chọn", "Hãy chọn template.", parent=self)
            return
        try:
            self.svc.set_default(tid)
            self._refresh_list()
            self.tree.selection_set(str(tid))
        except Exception as e:
            messagebox.showerror("Lỗi", str(e), parent=self)

    def _save_name(self):
        tid = self._get_selected_id()
        if not tid:
            return
        try:
            self.svc.update_name(tid, self.var_ten.get(), self.var_mo_ta.get())
            self._refresh_list()
            self.tree.selection_set(str(tid))
        except Exception as e:
            messagebox.showerror("Lỗi", str(e), parent=self)

    def _validate(self):
        tpl = self._get_selected_tpl()
        if not tpl or not tpl.get('file_path'):
            messagebox.showwarning("Chưa chọn", "Hãy chọn template có file.", parent=self)
            return
        if not os.path.exists(tpl['file_path']):
            messagebox.showerror("Lỗi", f"File không tồn tại:\n{tpl['file_path']}", parent=self)
            return

        result = self.svc.validate_template_file(tpl['file_path'])

        self.txt_validate.config(state='normal')
        self.txt_validate.delete('1.0', 'end')

        if result.get('valid'):
            self.txt_validate.insert('end', f"✅ Template hợp lệ — {result.get('total_keys', 0)} placeholder\n\n")
        else:
            self.txt_validate.insert('end', f"⚠ Có {len(result.get('errors', []))} vấn đề:\n\n")
            for err in result.get('errors', []):
                self.txt_validate.insert('end', f"  ❌ {err}\n")

        self.txt_validate.insert('end', "\nCác placeholder tìm thấy:\n")
        for ph in result.get('placeholders', []):
            mark = "✅" if ph['valid'] else "❓"
            self.txt_validate.insert('end', f"  {mark} {{{{{ph['key']}}}}}\n")

        self.txt_validate.config(state='disabled')

    def _preview_export(self):
        if not self.hp_id:
            messagebox.showwarning("Chưa chọn HP", "Không có học phần nào đang mở.", parent=self)
            return
        tpl = self._get_selected_tpl()
        if not tpl or not tpl.get('file_path'):
            messagebox.showwarning("Chưa chọn", "Hãy chọn template.", parent=self)
            return

        out = filedialog.asksaveasfilename(
            defaultextension='.docx',
            filetypes=[("Word", "*.docx")],
            title="Lưu file preview"
        )
        if not out:
            return

        try:
            self.svc.export_with_template(self.hp_id, tpl['id'], out)
            if messagebox.askyesno("Thành công", f"Đã xuất thành công!\n{out}\n\nMở file ngay?", parent=self):
                os.startfile(out)
        except Exception as e:
            messagebox.showerror("Lỗi xuất", str(e), parent=self)

    def _export_with_this(self):
        tid = self._get_selected_id()
        if not tid:
            messagebox.showwarning("Chưa chọn", "Hãy chọn template.", parent=self)
            return
        self.selected_template_id = tid
        self.destroy()

    def _show_placeholders(self):
        placeholders = self.svc.get_placeholders_help(self.hp_id)
        # Kiểm tra xem có trường mở rộng nào được phát hiện không
        new_fields = [p for p in placeholders if 'Phát hiện từ dữ liệu' in p.get('group', '')]
        if new_fields:
            messagebox.showinfo("Đồng bộ Placeholder", 
                               f"Hệ thống đã tự động phát hiện và bổ sung {len(new_fields)} trường dữ liệu mở rộng "
                               "từ học phần này vào danh sách tra cứu.", parent=self)
        
        PlaceholderHelpDialog(self, placeholders)

    def _clear_detail(self):
        self.var_ten.set('')
        self.var_mo_ta.set('')
        self.var_path.set('')


# ── Helper Dialogs ────────────────────────────────────────────────────────────

class _TextInputDialog(tb.Toplevel):
    def __init__(self, parent, title, prompt, default=''):
        super().__init__(parent)
        self.title(title)
        self.geometry("380x140")
        self.transient(parent)
        self.grab_set()
        self.result = None
        self.resizable(False, False)

        tb.Label(self, text=prompt, font=('Segoe UI', 10)).pack(padx=16, pady=(16, 4), anchor='w')
        self.var = tk.StringVar(value=default)
        entry = tb.Entry(self, textvariable=self.var, width=40)
        entry.pack(padx=16, fill='x')
        entry.focus_set()
        entry.select_range(0, 'end')

        btn_f = tb.Frame(self)
        btn_f.pack(pady=10)
        tb.Button(btn_f, text="OK", bootstyle='success', width=8,
                  command=self._ok).pack(side='left', padx=4)
        tb.Button(btn_f, text="Hủy", bootstyle='secondary-outline', width=8,
                  command=self.destroy).pack(side='left', padx=4)
        self.bind('<Return>', lambda _: self._ok())
        self.bind('<Escape>', lambda _: self.destroy())

    def _ok(self):
        self.result = self.var.get().strip()
        self.destroy()


class PlaceholderHelpDialog(tb.Toplevel):
    """Dialog hiển thị danh sách placeholder + cú pháp dùng trong template Word."""

    def __init__(self, parent, placeholders: list):
        super().__init__(parent)
        self.title("📋 Hướng dẫn Placeholder Template")
        self.geometry("700x550")
        self.transient(parent)
        self.grab_set()
        self.resizable(True, True)

        tb.Label(self, text="  Danh sách Placeholder có thể dùng trong file .docx template",
                 font=('Segoe UI', 11, 'bold'), bootstyle='inverse-info').pack(
                     fill='x', pady=(0, 8), ipady=6)

        # Treeview
        self.tree = tb.Treeview(
            self, columns=('key', 'type', 'group', 'example'),
            show='headings', bootstyle='info'
        )
        self.tree.heading('key', text='Placeholder')
        self.tree.heading('type', text='Kiểu')
        self.tree.heading('group', text='Nhóm')
        self.tree.heading('example', text='Cách dùng')
        self.tree.column('key', width=160, anchor='w')
        self.tree.column('type', width=130, anchor='w')
        self.tree.column('group', width=120, anchor='center')
        self.tree.column('example', width=230, anchor='w')

        sb = tb.Scrollbar(self, orient='vertical', command=self.tree.yview)
        self.tree.configure(yscrollcommand=sb.set)
        self.tree.pack(side='left', fill='both', expand=True, padx=(8, 0), pady=8)
        sb.pack(side='right', fill='y', pady=8, padx=(0, 8))

        groups = {}
        for ph in placeholders:
            g = ph.get('group', 'Khác')
            groups.setdefault(g, []).append(ph)

        for group, items in groups.items():
            gid = self.tree.insert('', 'end', values=(f'── {group} ──', '', '', ''), tags=('group',))
            for ph in items:
                self.tree.insert(gid, 'end', values=(
                    '{{ ' + ph['key'] + ' }}',
                    ph.get('type', ''),
                    ph.get('group', ''),
                    ph.get('example', '')
                ))
            self.tree.item(gid, open=True)

        self.tree.tag_configure('group', font=('Segoe UI', 9, 'bold'), foreground='#89b4fa')

        # Click to copy
        def _copy_key(event):
            item = self.tree.identify_row(event.y)
            if not item: return
            vals = self.tree.item(item, 'values')
            if vals and vals[0] and not vals[0].startswith('──'):
                key = vals[0]
                self.clipboard_clear()
                self.clipboard_append(key)
                # Show subtle toast or status
                old_title = self.title()
                self.title(f"✅ Đã chép: {key}")
                self.after(1000, lambda: self.title(old_title))

        self.tree.bind('<Double-1>', _copy_key)
        self.tree.bind('<Return>', _copy_key)

        # Action Buttons
        btn_bar = tb.Frame(self)
        btn_bar.pack(fill='x', padx=8, pady=4)
        
        tb.Label(btn_bar, text="💡 Nháy đúp vào placeholder để sao chép nhanh", 
                 font=('Segoe UI', 8, 'italic'), bootstyle='muted').pack(side='left')

        tb.Button(btn_bar, text="📥 Xuất danh sách (Excel)", bootstyle='info-outline',
                  command=self._export_list).pack(side='right', padx=2)

        # Cú pháp nhanh
        note = tb.Frame(self)
        note.pack(fill='x', padx=8, pady=(0, 8))
        tb.Label(note, text="💡 Cú pháp Jinja2 trong Word (gõ trực tiếp vào ô trong .docx):",
                 font=('Segoe UI', 9, 'bold')).pack(anchor='w')
        tb.Label(note, text="  {{ CourseName }}  |  {% for clo in CLOs %}  {{ clo.Code }}: {{ clo.Desc }}  {% endfor %}",
                 font=('Consolas', 9), foreground='#a6e3a1').pack(anchor='w')

        tb.Button(self, text="✖ Đóng", bootstyle='secondary-outline', command=self.destroy).pack(pady=4)

    def _export_list(self):
        """Xuất danh sách placeholder ra Excel để người dùng in hoặc tra cứu."""
        import openpyxl
        from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
        from tkinter import filedialog
        
        out = filedialog.asksaveasfilename(
            defaultextension='.xlsx',
            filetypes=[("Excel", "*.xlsx")],
            title="Xuất danh sách Placeholder"
        )
        if not out: return
        
        try:
            # Tạo workbook
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Placeholders"

            # Styles
            header_font = Font(bold=True, color="FFFFFF")
            header_fill = PatternFill("solid", fgColor="1E3A5F")
            center_align = Alignment(horizontal="center", vertical="center")
            border = Border(left=Side(style='thin'), right=Side(style='thin'), 
                            top=Side(style='thin'), bottom=Side(style='thin'))

            headers = ['Placeholder', 'Kiểu dữ liệu', 'Nhóm', 'Ví dụ cách dùng']
            for col, text in enumerate(headers, 1):
                cell = ws.cell(row=1, column=col, value=text)
                cell.font = header_font
                cell.fill = header_fill
                cell.alignment = center_align
                cell.border = border

            # Thu thập dữ liệu từ treeview (đệ quy để lấy hết item trong nhóm)
            def get_all_rows(parent=''):
                data_rows = []
                for item_id in self.tree.get_children(parent):
                    vals = self.tree.item(item_id, 'values')
                    # Bỏ qua dòng tiêu đề nhóm (bắt đầu bằng ──)
                    if vals and vals[0] and not vals[0].startswith('──'):
                        data_rows.append(vals)
                    # Đệ quy lấy con của group
                    data_rows.extend(get_all_rows(item_id))
                return data_rows

            rows = get_all_rows()
            
            for r_idx, row_vals in enumerate(rows, 2):
                for c_idx, val in enumerate(row_vals, 1):
                    cell = ws.cell(row=r_idx, column=c_idx, value=val)
                    cell.border = border
                    if c_idx == 3: # Cột Nhóm căn giữa
                         cell.alignment = center_align
            
            # Cân chỉnh độ rộng cột
            ws.column_dimensions['A'].width = 25
            ws.column_dimensions['B'].width = 20
            ws.column_dimensions['C'].width = 15
            ws.column_dimensions['D'].width = 45

            wb.save(out)
            messagebox.showinfo("Thành công", f"Đã xuất danh sách thành công!\n{out}", parent=self)
        except Exception as e:
            messagebox.showerror("Lỗi", f"Không thể xuất file: {e}", parent=self)
